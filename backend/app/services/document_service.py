"""
Document text extraction — pdf, docx, txt, md.

Used by the admin file-upload ingestion endpoint.
All extraction runs in a thread pool (CPU-bound, blocking libraries).
"""

from __future__ import annotations

import asyncio
import io

from app.core.logging import get_logger

log = get_logger("documents")

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileType(Exception):
    pass


class FileTooLarge(Exception):
    pass


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(io.BytesIO(data))
    pages  = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document  # type: ignore

    doc        = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract table cell text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n\n".join(paragraphs)


def _extract_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _get_extension(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx != -1 else ""


async def extract_text(filename: str, data: bytes) -> str:
    """
    Extract plain text from an uploaded file.

    Raises:
        UnsupportedFileType — extension not in SUPPORTED_EXTENSIONS
        FileTooLarge        — file exceeds MAX_FILE_SIZE
        ValueError          — extraction produced no usable text
    """
    if len(data) > MAX_FILE_SIZE:
        raise FileTooLarge(
            f"File exceeds maximum size of {MAX_FILE_SIZE // (1024*1024)}MB "
            f"(got {len(data) / (1024*1024):.1f}MB)"
        )

    ext = _get_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    loop = asyncio.get_event_loop()

    if ext == ".pdf":
        text = await loop.run_in_executor(None, _extract_pdf, data)
    elif ext == ".docx":
        text = await loop.run_in_executor(None, _extract_docx, data)
    else:  # .txt, .md
        text = await loop.run_in_executor(None, _extract_text, data)

    text = text.strip()
    if not text or len(text) < 10:
        raise ValueError("Document contains no extractable text")

    log.info("document_extracted", filename=filename, ext=ext, chars=len(text))
    return text
